import streamlit as st
import requests
import re
from difflib import SequenceMatcher
from bs4 import BeautifulSoup
from html import escape
from docx import Document
from io import BytesIO

# Configuración de la página
st.set_page_config(page_title="Corrector de DOCX con Resaltado de Cambios", layout="wide")

# Título de la aplicación
st.title("Corrector de DOCX con Resaltado de Cambios")

# Instrucciones para el usuario
st.markdown("""
**Instrucciones:**
1. **Sube un archivo DOCX** que contenga el texto a corregir. Asegúrate de que las notas a pie de página estén en el formato `[1]`, `[2]`, etc., y que las citas textuales estén entre comillas dobles `"cita"`.
2. La aplicación corregirá la ortografía, gramática y estilo del texto principal.
3. **Las notas a pie de página y citas textuales no serán modificadas.**
4. Los cambios se resaltarán en colores:
   - 🟨 **Amarillo**: Cambios realizados.
   - 🟥 **Rojo**: Eliminaciones.
   - 🟩 **Verde**: Adiciones.
""")

# Función para extraer texto del DOCX y convertir superíndices en [n]
def extract_text_from_docx(file):
    doc = Document(file)
    full_text = []
    footnote_map = {}  # Mapa para almacenar notas a pie de página si es necesario
    footnote_counter = 1  # Contador de notas a pie

    for para in doc.paragraphs:
        para_text = ""
        for run in para.runs:
            if run.font.superscript:
                # Detectar si el texto del run es un número
                if run.text.isdigit():
                    # Reemplazar superíndice con [n]
                    para_text += f"[{run.text}]"
                else:
                    para_text += run.text
            else:
                para_text += run.text
        full_text.append(para_text)
    return '\n'.join(full_text)

# Función para contar palabras
def count_words(text):
    return len(re.findall(r'\b\w+\b', text))

# Función para proteger notas a pie de página y citas textuales
def protect_text(text):
    placeholders = {}
    protected_text = text

    # Proteger citas textuales (texto entre comillas dobles)
    quotes = re.findall(r'"(.*?)"', protected_text)
    for idx, quote in enumerate(quotes):
        placeholder = f"__QUOTE_{idx}__"
        placeholders[placeholder] = f'"{quote}"'
        protected_text = protected_text.replace(f'"{quote}"', placeholder)

    # Proteger notas a pie de página en formato [1], [2], etc.
    footnotes = re.findall(r'\[\d+\]', protected_text)
    for idx, footnote in enumerate(footnotes):
        placeholder = f"__FOOTNOTE_{idx}__"
        placeholders[placeholder] = footnote
        protected_text = protected_text.replace(footnote, placeholder)

    return protected_text, placeholders

# Función para restaurar notas a pie de página y citas textuales
def restore_text(text, placeholders):
    for placeholder, original in placeholders.items():
        text = text.replace(placeholder, original)
    return text

# Función para llamar a la API de Together
def correct_text(text, api_key):
    url = "https://api.together.xyz/v1/chat/completions"
    headers = {
        "Authorization": f"Bearer {api_key}",
        "Content-Type": "application/json"
    }

    # Mensaje de sistema para la API
    system_message = (
        "Eres un asistente que corrige la ortografía, gramática y estilo de un texto. "
        "No debes modificar las notas a pie de página ni las citas textuales (texto entre comillas)."
    )

    # Mensaje de usuario con el texto a corregir
    user_message = text

    payload = {
        "model": "mistralai/Mixtral-8x7B-Instruct-v0.1",
        "messages": [
            {"role": "system", "content": system_message},
            {"role": "user", "content": user_message}
        ],
        "max_tokens": 2512,
        "temperature": 0.7,
        "top_p": 0.7,
        "top_k": 50,
        "repetition_penalty": 1,
        "stop": ["<|eot_id|>"],
        "stream": False  # Cambiado a False para simplificar
    }

    try:
        response = requests.post(url, headers=headers, json=payload)
        response.raise_for_status()
        data = response.json()
        # Asumiendo que la respuesta contiene el texto corregido en 'choices'[0]['message']['content']
        corrected_text = data.get("choices", [{}])[0].get("message", {}).get("content", "")
        return corrected_text
    except requests.exceptions.HTTPError as http_err:
        st.error(f"Error en la API: {response.status_code} - {response.text}")
        return None
    except Exception as err:
        st.error(f"Error inesperado: {err}")
        return None

# Función para resaltar los cambios
def highlight_changes(original, corrected):
    matcher = SequenceMatcher(None, original, corrected)
    highlighted = ""
    for tag, i1, i2, j1, j2 in matcher.get_opcodes():
        if tag == 'equal':
            highlighted += escape(original[i1:i2])
        elif tag == 'replace':
            highlighted += f"<span style='background-color: yellow;'>{escape(corrected[j1:j2])}</span>"
        elif tag == 'delete':
            highlighted += f"<span style='background-color: red; text-decoration: line-through;'>{escape(original[i1:i2])}</span>"
        elif tag == 'insert':
            highlighted += f"<span style='background-color: green;'>{escape(corrected[j1:j2])}</span>"
    return highlighted

# Área para subir el archivo DOCX
uploaded_file = st.file_uploader("Sube tu archivo DOCX aquí:", type=["docx"])

if uploaded_file is not None:
    # Extraer el texto del archivo DOCX y convertir superíndices en [n]
    text = extract_text_from_docx(uploaded_file)

    st.markdown("### Texto Original")
    st.write(text)

    # Botón para iniciar la corrección
    if st.button("Corregir Texto"):
        if not text.strip():
            st.warning("El archivo subido está vacío.")
        else:
            total_words = count_words(text)
            if total_words > 2000:
                st.error(f"El texto excede el límite de 2000 palabras. Actualmente tiene {total_words} palabras.")
            else:
                with st.spinner("Corrigiendo el texto..."):
                    # Proteger las notas a pie de página y las citas textuales
                    protected_text, placeholders = protect_text(text)

                    # Llamar a la API con el texto protegido
                    api_key = st.secrets["together_api_key"]
                    corrected_protected = correct_text(protected_text, api_key)

                    if corrected_protected:
                        # Restaurar las notas a pie de página y las citas textuales
                        corrected = restore_text(corrected_protected, placeholders)

                        # Resaltar los cambios
                        highlighted_text = highlight_changes(text, corrected)

                        # Dividir la página en dos columnas
                        col1, col2 = st.columns(2)

                        with col1:
                            st.header("Texto Original")
                            st.write(text)

                        with col2:
                            st.header("Texto Corregido")
                            st.markdown(f"<div style='white-space: pre-wrap;'>{highlighted_text}</div>", unsafe_allow_html=True)
